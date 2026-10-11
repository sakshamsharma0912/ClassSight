"""Faculty-generated department documents, with stable warning snapshots."""
import csv
import io
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                               Spacer, Table, TableStyle)

from portal_core import COMPONENTS, points, safe_cell, summarize

NAVY = colors.HexColor("#153457")
TEAL = colors.HexColor("#047c85")
LIGHT = colors.HexColor("#edf3f8")
INK = colors.HexColor("#1e334d")
RED = colors.HexColor("#b42336")
FONT_DIR = Path(__file__).parent / "assets" / "fonts"
for name, filename in (("Portal", "DejaVuSans.ttf"), ("Portal-Bold", "DejaVuSans-Bold.ttf")):
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(FONT_DIR / filename)))


def p(text, size=8, bold=False, color=INK, alignment=TA_LEFT):
    return Paragraph(escape(str(text)).replace("\n", "<br/>"), ParagraphStyle(
        "cell", fontName="Portal-Bold" if bold else "Portal", fontSize=size,
        leading=size * 1.4, textColor=color, alignment=alignment, spaceAfter=0))


def pct(value):
    return "—" if value is None else f"{value:.2f}%"


def scope_text(filters):
    return " | ".join([filters.get("section") or "All sections", filters.get("batch") or "All batches",
                       filters.get("course") or "All courses"] +
                      ([f"Search: {filters['q']}"] if filters.get("q") else []) +
                      ([f"Attendance filter: {filters['status']}"] if filters.get("status") else []) +
                      (["Individual student record"] if filters.get("student") else []))


def header(profile, title, generated, width):
    content = [
        p(profile["institution"], 14, True, NAVY),
        p(profile["campus"] + " • " + profile["department"], 9),
        Spacer(1, 7), p(title, 12, True, TEAL), Spacer(1, 4),
        p(f"Academic year {profile['academic_year']} | Semester {profile['semester']} | "
          f"Prepared by {profile['name']}" + (f" ({profile['faculty_no']})" if profile["faculty_no"] else ""), 8),
        p("Generated " + generated.strftime("%d %b %Y, %H:%M IST"), 8),
        Spacer(1, 10),
    ]
    return content


def footer(canvas, doc):
    width, _height = doc.pagesize
    canvas.setStrokeColor(colors.HexColor("#cdd8e3"))
    canvas.line(doc.leftMargin, 31, width - doc.rightMargin, 31)
    canvas.setFont("Portal", 7)
    canvas.setFillColor(colors.HexColor("#657589"))
    canvas.drawString(doc.leftMargin, 20, "ClassSight | Faculty-maintained academic record")
    canvas.drawRightString(width - doc.rightMargin, 20, "Page " + str(doc.page))


def table(data, widths, font_size=7.1, repeats=1):
    rows = [[p(cell, font_size, bold=(index == 0),
               color=colors.white if index == 0 else INK) for cell in row]
            for index, row in enumerate(data)]
    result = Table(rows, colWidths=widths, repeatRows=repeats, hAlign="LEFT")
    result.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, 0), (-1, 0), 1, NAVY),
        ("LINEBELOW", (0, 1), (-1, -1), 0.35, colors.HexColor("#dce4ed")),
    ]))
    return result


def attendance_headers():
    return ["Registration no.", "Student", "Section", "Batch", "Course", "Max hours",
            "Attended", "Absent", "Attendance %", "Margin (h)", "Attendance status"]


def attendance_values(r):
    return [r["student_id"], r["name"], r["section"], r["batch"], r["course_code"],
            r["max_hours"], r["attended_hours"], r["absent_hours"], pct(r["attendance_pct"]),
            f"{r['margin']:+d}" if r["max_hours"] else "—", r["attendance_status"]]


def marks_headers(short=False):
    labels = (["Assignment 1\n/5", "Assignment 2\n/5", "Internal Exam 1\n/15",
               "Mini Project\n/10", "Practical\n/10", "Internal Exam 2\n/15",
               "University Final Exam\n/40"] if short else
              [label + " /" + str(cap) for _key, label, cap in COMPONENTS])
    return ["Registration no.", "Student", "Section", "Course"] + labels + ["Internal /60", "Total /100", "Grade"]


def marks_values(r):
    return [r["student_id"], r["name"], r["section"], r["course_code"],
            *(points(r[key]) if r[key] is not None else "Pending" for key, _label, _cap in COMPONENTS),
            f"{r['internal']:g}" + ("" if r["internal_complete"] else " *"),
            f"{r['total']:g}" if r["complete"] else "Pending", r["grade"]]


def department_pdf(profile, rows, filters, generated, personal=False):
    buffer = io.BytesIO()
    title = "Student Academic Performance Report" if personal else "Department Academic Performance Report"
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=30, leftMargin=30,
                            topMargin=27, bottomMargin=43, title="ClassSight - " + title,
                            author=profile["name"])
    width = landscape(A4)[0] - 60
    story = header(profile, title, generated, width)
    story += [p(scope_text(filters), 8), Spacer(1, 10)]
    summary = summarize(rows)
    cards = [["Students", "Course records", "Attendance", "Shortage records", "Evaluated average"],
             [summary["students"], summary["records"], pct(summary["attendance_pct"]), summary["at_risk"],
              "Pending" if summary["average"] is None else f"{summary['average']:.2f}/100"]]
    story += [table(cards, [width / 5] * 5, 8), Spacer(1, 13),
              p("Course-wise attendance register", 10, True, NAVY), Spacer(1, 7)]
    attendance = [attendance_headers()] + [attendance_values(r) for r in rows]
    attendance.append(["Student summary" if personal else "Class summary", f"{summary['students']} students", "", "", "",
                       summary["max_hours"], summary["attended_hours"], summary["absent_hours"],
                       pct(summary["attendance_pct"]), f"{summary['margin']:+d}",
                       f"{summary['at_risk']} below 75%"])
    widths = [86, 112, 45, 59, 75, 45, 48, 43, 65, 53, 90]
    attendance_table = table(attendance, [w * width / sum(widths) for w in widths], 6.7)
    attendance_table.setStyle(TableStyle([("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#dbeaf1"))]))
    for i, row in enumerate(rows, 1):
        if row["attendance_status"] == "Below 75%":
            attendance_table.setStyle(TableStyle([("BACKGROUND", (8, i), (-1, i), colors.HexColor("#fbe9ec"))]))
    story += [attendance_table, Spacer(1, 9),
              p("Max hours are conducted/recorded hours. Summary hours are student-hours; the percentage is weighted. "
                "Margin = attended - ceil(0.75 × max hours). Eligibility is checked individually for every course; "
                "the aggregate does not override a shortage. No condonation is applied automatically.", 7)]
    if not rows:
        story += [Spacer(1, 8), p("No academic records match the selected filters.", 9)]
    story += [PageBreak()]
    story += header(profile, "Component-wise Mark Evaluation", generated, width)
    story += [p(scope_text(filters), 8), Spacer(1, 10)]
    data = [marks_headers(True)] + [marks_values(r) for r in rows]
    # Scale a dense 14-column grid to the landscape page without clipping.
    widths = [80, 95, 40, 65] + [55, 55, 52, 49, 47, 52, 75, 47, 51, 45]
    scale = width / sum(widths)
    story += [table(data, [w * scale for w in widths], 6.3), Spacer(1, 10),
              p("Blank marks remain Pending. * indicates a partial internal total. The configured project scheme is "
                "5 + 5 + 15 + 10 + 10 + 15 + 40 = 100 marks, in the displayed order. "
                "Grade bands: O ≥90; A+ ≥80; A ≥70; B+ ≥60; B ≥50; C ≥40; F <40. "
                "Marks and attendance eligibility are reported separately.", 7),
              Spacer(1, 24)]
    signatures = [["Prepared by", "Reviewed by / Department Head"],
                  [profile["name"], "Signature: __________________________"],
                  ["Signature: __________________________", "Date: __________________________"]]
    story += [table(signatures, [width / 2] * 2, 8)]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def report_csv(rows, personal=False):
    headers = attendance_headers() + [label + " /" + str(cap) for _key, label, cap in COMPONENTS] + [
        "Internal /60", "Evaluation", "Total /100", "Grade", "Can miss hours", "Recovery hours"]
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([safe_cell(v) for v in attendance_values(row) +
                         [points(row[k]) for k, _label, _cap in COMPONENTS] +
                         [row["internal"], "Complete" if row["complete"] else "Pending",
                          row["total"] if row["complete"] else "", row["grade"],
                          row["can_miss"], row["recover_hours"]]])
    summary = summarize(rows)
    writer.writerow(["Student summary" if personal else "Class summary", f"{summary['students']} students", "", "", "",
                     summary["max_hours"], summary["attended_hours"], summary["absent_hours"],
                     pct(summary["attendance_pct"]), summary["margin"], f"{summary['at_risk']} below 75%"])
    return stream.getvalue().encode("utf-8-sig")


def department_workbook(profile, rows, filters, generated, personal=False):
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = summarize(rows)
    meta = workbook.create_sheet("Report")
    meta.append([profile["institution"]])
    meta.append([profile["campus"], profile["department"]])
    meta.append(["ClassSight - " + ("Student Academic Performance Report" if personal else "Department Academic Performance Report")])
    meta.append(["Academic year", profile["academic_year"], "Semester", profile["semester"]])
    meta.append(["Prepared by", profile["name"], "Faculty ID", profile["faculty_no"]])
    meta.append(["Generated", generated.strftime("%d %b %Y %H:%M IST")])
    meta.append(["Scope", scope_text(filters)])
    meta.append(["Students", summary["students"]])
    meta.append(["Course records", summary["records"]])
    meta.append(["Max student-hours", summary["max_hours"]])
    meta.append(["Attended student-hours", summary["attended_hours"]])
    meta.append(["Absent student-hours", summary["absent_hours"]])
    meta.append(["Weighted attendance %", summary["attendance_pct"]])
    meta.append(["Records below 75%", summary["at_risk"]])
    meta.append(["Completed evaluations", summary["complete_count"]])
    meta.append(["Evaluated average /100", summary["average"]])
    meta.append(["Attendance rule", "75% in each course; aggregate does not override course shortage."])
    meta.append(["Margin formula", "Attended - ceil(0.75 × max hours)."])
    meta.append(["Mark scheme", "5, 5, 15, 10, 10, 15, 40; blank marks are Pending."])
    meta.append(["Origin", "Faculty-maintained academic record from ClassSight."])
    attendance = workbook.create_sheet("Attendance")
    attendance.append(attendance_headers() + ["Can miss hours", "Recovery hours"])
    for row in rows:
        attendance.append([row["student_id"], row["name"], row["section"], row["batch"], row["course_code"],
                           row["max_hours"], row["attended_hours"], row["absent_hours"],
                           row["attendance_pct"], row["margin"] if row["max_hours"] else None,
                           row["attendance_status"], row["can_miss"], row["recover_hours"]])
    attendance.append(["Student summary" if personal else "Class summary", f"{summary['students']} students", "", "", "",
                       summary["max_hours"], summary["attended_hours"], summary["absent_hours"],
                       summary["attendance_pct"], summary["margin"], f"{summary['at_risk']} below 75%"])
    marks = workbook.create_sheet("Marks")
    marks.append(marks_headers())
    for row in rows:
        marks.append([row["student_id"], row["name"], row["section"], row["course_code"],
                      *(row[k] / 100 if row[k] is not None else None for k, _label, _cap in COMPONENTS),
                      row["internal"] if row["internal_complete"] else None,
                      row["total"] if row["complete"] else None, row["grade"]])
    border = Side(style="thin", color="DEE7F0")
    for sheet in workbook:
        sheet.freeze_panes = "B2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions if sheet.title != "Report" else "A1:D1"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = "1:1"
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = safe_cell(cell.value)
                cell.font = Font(name="Calibri", size=11, color="213750")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=border)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="153457")
                    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F5FA")
            sheet.row_dimensions[row[0].row].height = 32 if row[0].row == 1 else 29
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 18
        sheet.column_dimensions["B"].width = 25 if sheet.title != "Report" else 75
    for row in attendance.iter_rows(min_row=2, max_row=attendance.max_row - 1):
        if row[10].value == "Below 75%":
            for cell in row[8:11]:
                cell.fill = PatternFill("solid", fgColor="FCE8ED")
                cell.font = Font(color="AE243B")
    for cell in attendance["I"][1:]:
        cell.number_format = '0.00"%"'
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def warning_pdf(snapshot, reference, issued_at):
    profile, student, row = snapshot["profile"], snapshot["student"], snapshot["record"]
    generated = datetime.fromisoformat(issued_at)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=42, rightMargin=42,
                            topMargin=37, bottomMargin=45, title="Attendance Shortage Notice",
                            author=profile["name"])
    width = A4[0] - 84
    story = header(profile, "Attendance Shortage Notice", generated, width)
    story += [p("Reference: " + reference, 8, True), Spacer(1, 15),
              p("To: " + student["name"], 10, True),
              p("Registration no.: " + student["student_id"] + "\nSection: " + student["section"] +
                " | Batch: " + student["batch"], 9),
              Spacer(1, 15),
              p("Subject: Attendance below the 75% requirement", 10, True, RED), Spacer(1, 12),
              p("The current attendance record for the following course is below the minimum 75% "
                "requirement. Please meet the course faculty and academic advisor promptly to review "
                "the shortage and plan attendance for the remaining classes.", 10), Spacer(1, 13)]
    detail = [["Course", row["course_code"] + " - " + row["course_title"]],
              ["Max / conducted hours", row["max_hours"]],
              ["Attended hours", row["attended_hours"]], ["Absent hours", row["absent_hours"]],
              ["Attendance", pct(row["attendance_pct"])],
              ["Required at current hours", row["required_hours"]],
              ["Current margin", f"{row['margin']:+d} hours"],
              ["Consecutive hours to recover 75%", row["recover_hours"]]]
    story += [table([["Attendance record", "Value"]] + detail, [205, width - 205], 9),
              Spacer(1, 15),
              p("Recovery hours assume that every additional hour is attended and no further hours are "
                "missed. If fewer classes remain, recovery may not be possible.", 9),
              Spacer(1, 11),
              p("Attendance eligibility is determined for each course separately. This notice records "
                "the current shortage for faculty review; final examination eligibility and any approved "
                "condonation must be confirmed by the institution.", 9),
              Spacer(1, 26),
              p("Course faculty / Academic advisor", 10, True),
              p(profile["name"] + (f" | Faculty ID: {profile['faculty_no']}" if profile["faculty_no"] else ""), 9),
              Spacer(1, 10), p("Signature: __________________________       Date: __________________", 9)]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
