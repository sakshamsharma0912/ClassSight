import csv
import io
import json
import re
import sqlite3
from datetime import date
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from werkzeug.security import check_password_hash, generate_password_hash

from app import create_app, get_db
from portal_core import (COMPONENTS, COURSES, CSV_FIELDS, attendance_metrics,
                         enrich, grade_for, mark, points, summarize)


@pytest.fixture
def app(tmp_path):
    return create_app({"TESTING": True, "DATABASE": str(tmp_path / "faculty.sqlite3"),
                       "SECRET_KEY": "test-key-" * 8})


def token(client):
    with client.session_transaction() as session:
        return session.get("_csrf_token")


def register(client, email="faculty@example.edu", name="Course Faculty"):
    client.get("/register")
    return client.post("/register", data={"csrf_token": token(client), "name": name, "email": email,
                                         "password": "FacultyPass123", "confirm_password": "FacultyPass123"})


@pytest.fixture
def client(app):
    client = app.test_client()
    assert register(client).status_code == 302
    client.get("/teacher/dashboard")
    return client


def post(client, path, data=None, **kwargs):
    values = dict(data or {})
    values["csrf_token"] = token(client)
    return client.post(path, data=values, **kwargs)


def rows(app, sql="SELECT * FROM records ORDER BY id", params=()):
    with app.app_context():
        return [dict(r) for r in get_db().execute(sql, params)]


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links, self.assets, self.forms, self.dialogs, self.triggers = [], [], [], set(), set()
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "a":
            self.links.append(attrs.get("href", ""))
        if tag in ("script", "link"):
            self.assets.append(attrs.get("src") or attrs.get("href", ""))
        if tag == "form":
            self.forms.append(attrs)
        if tag == "dialog":
            self.dialogs.add(attrs.get("id"))
        if "data-dialog" in attrs:
            self.triggers.add(attrs["data-dialog"])


def test_authentication_hash_and_csrf(app, client):
    faculty = rows(app, "SELECT * FROM teachers")[0]
    assert faculty["password_hash"].startswith("scrypt:")
    assert check_password_hash(faculty["password_hash"], "FacultyPass123")
    assert client.post("/students/1/delete").status_code == 400
    assert post(client, "/logout").status_code == 302
    for path in ("/students", "/attendance", "/marks", "/reports/pdf", "/warnings", "/csv/export"):
        assert client.get(path).location.endswith("/login")
    client.get("/login")
    bad = post(client, "/login", {"email": faculty["email"], "password": "incorrect"})
    assert bad.status_code == 400
    good = post(client, "/login", {"email": faculty["email"], "password": "FacultyPass123"})
    assert good.status_code == 302


def test_login_throttle(app):
    client = app.test_client()
    client.get("/login")
    for _ in range(10):
        assert post(client, "/login", {"email": "missing@example.edu", "password": "wrong"}).status_code == 400
    assert post(client, "/login", {"email": "missing@example.edu", "password": "wrong"}).status_code == 429


def test_seed_once_and_independent_accounts(app, client):
    assert len(rows(app, "SELECT * FROM students")) == 16
    assert len(rows(app)) == 112
    second = app.test_client()
    assert register(second, "second@example.edu").status_code == 302
    second.get("/teacher/dashboard")
    assert len(rows(app, "SELECT * FROM students")) == 32
    for path in ("/students/1", "/students/1/edit", "/records/1/marks", "/reports/pdf?student=1"):
        assert second.get(path).status_code == 404
    assert post(second, "/students/1/delete").status_code == 404
    assert post(second, "/warnings/issue/1").status_code == 404
    with app.app_context():
        db = get_db()
        with db:
            db.execute("DELETE FROM students")
    reopened = create_app(dict(app.config))
    assert rows(reopened, "SELECT * FROM students") == []


def test_security_headers_and_escaping(app, client):
    response = client.get("/teacher/dashboard")
    assert response.headers["X-Frame-Options"] == "DENY"
    assert "script-src 'self'" in response.headers["Content-Security-Policy"]
    assert response.headers["Cache-Control"] == "no-store"
    result = post(client, "/students/add", {"student_id": "XSS001", "name": "<script>alert(1)</script>",
                                          "section": "CSE-A", "batch": "2025-2029"})
    page = client.get(result.location).get_data(as_text=True)
    assert "<script>alert(1)</script>" not in page
    assert "&lt;script&gt;" in page
    cookie = response.headers.get("Set-Cookie", "")
    if cookie:
        assert "HttpOnly" in cookie and "SameSite=Lax" in cookie


@pytest.mark.parametrize("maximum,attended,margin,miss,recover,status", [
    (0, 0, 0, 0, 0, "Not started"), (4, 3, 0, 0, 0, "Eligible"),
    (100, 75, 0, 0, 0, "Eligible"), (100, 74, -1, 0, 4, "Below 75%"),
    (100, 90, 15, 20, 0, "Eligible"), (5, 3, -1, 0, 3, "Below 75%"),
    (7, 6, 0, 1, 0, "Eligible"), (10000, 7499, -1, 0, 4, "Below 75%"),
])
def test_attendance_boundaries(maximum, attended, margin, miss, recover, status):
    result = attendance_metrics(maximum, attended)
    assert (result["margin"], result["can_miss"], result["recover_hours"], result["attendance_status"]) == (
        margin, miss, recover, status)
    assert result["absent_hours"] == maximum - attended


def test_attendance_forward_formulas_exhaustively():
    for maximum in range(1, 101):
        for attended in range(maximum + 1):
            result = attendance_metrics(maximum, attended)
            if result["eligible"]:
                future = result["can_miss"]
                assert 4 * attended >= 3 * (maximum + future)
                assert 4 * attended < 3 * (maximum + future + 1)
            else:
                future = result["recover_hours"]
                assert 4 * (attended + future) >= 3 * (maximum + future)
                assert 4 * (attended + future - 1) < 3 * (maximum + future - 1)


def test_student_crud_and_section_batch_search(app, client):
    result = post(client, "/students/add", {"student_id": "CSE25D101", "name": "New Enrollee",
                                          "section": "CSE-D", "batch": "2025-2029"})
    assert result.status_code == 302
    pk = int(result.location.rsplit("/", 1)[1])
    records = rows(app, "SELECT * FROM records WHERE student_pk = ?", (pk,))
    assert len(records) == 7
    assert all(r["max_hours"] == 0 and r["university_final"] is None for r in records)
    scoped = client.get("/students?section=CSE-D&batch=2025-2029&q=New+Enrollee").get_data(as_text=True)
    assert "New Enrollee" in scoped and "Aarav Mehta" not in scoped
    result = post(client, f"/students/{pk}/edit", {"student_id": "CSE25D102", "name": "Updated Enrollee",
                                               "section": "CSE-B", "batch": "2024-2028", "email": "student@example.edu"})
    assert result.status_code == 302
    page = client.get(f"/students/{pk}").get_data(as_text=True)
    assert "Updated Enrollee" in page and "2024-2028" in page
    assert post(client, f"/students/{pk}/delete").status_code == 302
    assert client.get(f"/students/{pk}").status_code == 404
    assert rows(app, "SELECT * FROM records WHERE student_pk = ?", (pk,)) == []


@pytest.mark.parametrize("field,value", [("batch", "2029-2025"), ("section", "<script>"),
                                         ("student_id", "bad/id"), ("email", "invalid")])
def test_invalid_student_fields(client, field, value):
    profile = {"student_id": "NEW001", "name": "New Student", "section": "CSE-A", "batch": "2025-2029"}
    profile[field] = value
    assert post(client, "/students/add", profile).status_code == 400


def test_single_attendance_edit_and_stale_version(app, client):
    old = rows(app)[0]
    values = {"version": old["version"], "max_hours": "100", "attended_hours": "74"}
    assert post(client, "/records/1/attendance", values).status_code == 302
    updated = rows(app)[0]
    assert enrich(updated)["attendance_status"] == "Below 75%"
    assert post(client, "/records/1/attendance", values).status_code == 400
    bad = {"version": updated["version"], "max_hours": "10", "attended_hours": "11"}
    assert post(client, "/records/1/attendance", bad).status_code == 400
    assert rows(app)[0]["attended_hours"] == 74


def test_mark_caps_order_precision_and_pending(app, client):
    record = rows(app)[0]
    data = {"version": record["version"], **{key: str(cap) for key, _label, cap in COMPONENTS}}
    assert post(client, "/records/1/marks", data).status_code == 302
    result = enrich(rows(app)[0])
    assert result["total"] == 100 and result["internal"] == 60 and result["grade"] == "O"
    data["version"] = rows(app)[0]["version"]
    data["university_final"] = ""
    assert post(client, "/records/1/marks", data).status_code == 302
    assert enrich(rows(app)[0])["grade"] == "Pending"
    data["version"] = rows(app)[0]["version"]
    data["assignment_1"] = "5.01"
    assert post(client, "/records/1/marks", data).status_code == 400
    assert rows(app)[0]["assignment_1"] == 500
    assert mark("4.99", "Assignment", 5) == 499
    assert mark(0, "Assignment", 5) == 0 and points(499) == "4.99"
    assert grade_for(8999) == "A+" and grade_for(9000) == "O"
    page = client.get("/marks").get_data(as_text=True)
    header = page[page.index("<thead>"):page.index("</thead>")]
    positions = [header.index(label) for _key, label, _cap in COMPONENTS]
    assert positions == sorted(positions)


def grid_data(records, kind):
    data = {"record_id": [str(r["id"]) for r in records]}
    for row in records:
        prefix = "r" + str(row["id"]) + "_"
        data[prefix + "version"] = row["version"]
        if kind == "attendance":
            data[prefix + "max_hours"], data[prefix + "attended_hours"] = "100", "75"
        else:
            for key, _label, cap in COMPONENTS:
                data[prefix + key] = str(cap)
    return data


@pytest.mark.parametrize("kind,path", [("attendance", "/attendance/edit"), ("marks", "/marks")])
def test_bulk_register_atomic_save_and_stale_roster(app, client, kind, path):
    query = "?course=21CSC201J&section=CSE-A&batch=2025-2029"
    records = rows(app, """SELECT r.* FROM records r JOIN students s ON s.id=r.student_pk
                         WHERE s.section='CSE-A' AND r.course_code='21CSC201J'""")
    data = grid_data(records, kind)
    assert post(client, path + query, data).status_code == 302
    assert post(client, path + query, data).status_code == 400
    new = rows(app, """SELECT r.* FROM records r JOIN students s ON s.id=r.student_pk
                     WHERE s.section='CSE-A' AND r.course_code='21CSC201J'""")
    bad = grid_data(new, kind)
    bad["record_id"] = bad["record_id"][:-1]
    assert post(client, path + query, bad).status_code == 400
    invalid = grid_data(new, kind)
    key = "r" + str(new[-1]["id"]) + "_" + ("assignment_1" if kind == "marks" else "attended_hours")
    invalid[key] = "900"
    before = rows(app)
    assert post(client, path + query, invalid).status_code == 400
    assert rows(app) == before


def test_record_class_session_history_and_duplicate_submission(app, client):
    query = "?course=21CSC201J&section=CSE-A&batch=2025-2029"
    roster = rows(app, """SELECT r.* FROM records r JOIN students s ON s.id=r.student_pk
                       WHERE s.section='CSE-A' AND r.course_code='21CSC201J'""")
    data = {"record_id": [str(r["id"]) for r in roster], "hours": "2", "held_on": "2026-01-01"}
    for i, row in enumerate(roster):
        data["v" + str(row["id"])] = row["version"]
        data["a" + str(row["id"])] = "0" if i == 0 else "2"
    assert post(client, "/attendance/session" + query, data).status_code == 302
    first = rows(app, "SELECT * FROM records WHERE id = ?", (roster[0]["id"],))[0]
    assert first["max_hours"] == roster[0]["max_hours"] + 2
    assert first["attended_hours"] == roster[0]["attended_hours"]
    assert client.get("/attendance/sessions/1").status_code == 200
    assert len(rows(app, "SELECT * FROM attendance_sessions")) == 1
    assert post(client, "/attendance/session" + query, data).status_code == 400
    assert len(rows(app, "SELECT * FROM attendance_sessions")) == 1
    assert post(client, "/attendance/session", data).status_code == 400
    data["held_on"] = "2099-01-01"
    assert post(client, "/attendance/session" + query, data).status_code == 400


def csv_bytes(items):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(CSV_FIELDS)
    writer.writerows(items)
    return stream.getvalue().encode("utf-8-sig")


def upload(client, data, mode="reject", filename="import.csv"):
    return post(client, "/import", {"mode": mode, "file": (io.BytesIO(data), filename)},
                content_type="multipart/form-data")


def csv_row(reg="NEW001", name="Aditi Sharma", course="21CSC201J", attended=32, maximum=40):
    return (reg, name, "CSE-D", "2025-2029", course, maximum, attended, 4, 5, 12, 8, 9, 11, "")


def test_csv_import_multiple_courses_update_and_atomic_rollback(app, client):
    data = csv_bytes([csv_row(), csv_row(course="21CSC202J")])
    assert upload(client, data).status_code == 302
    added = rows(app, "SELECT * FROM students WHERE student_id='NEW001'")[0]
    assert len(rows(app, "SELECT * FROM records WHERE student_pk=?", (added["id"],))) == 7
    assert upload(client, data).status_code == 400
    assert upload(client, csv_bytes([csv_row(attended=35)]), "update").status_code == 302
    r = rows(app, "SELECT * FROM records WHERE student_pk=? AND course_code='21CSC201J'", (added["id"],))[0]
    assert r["attended_hours"] == 35 and r["university_final"] is None
    before = rows(app, "SELECT * FROM students")
    invalid = csv_bytes([csv_row(reg="NEW002"), csv_row(reg="NEW003", attended=41)])
    assert upload(client, invalid).status_code == 400
    assert rows(app, "SELECT * FROM students") == before
    # Conflict after the first INSERT also rolls back the whole transaction.
    conflict = csv_bytes([csv_row(reg="NEW004"), csv_row(name="Wrong Name")])
    assert upload(client, conflict, "update").status_code == 400
    assert rows(app, "SELECT * FROM students") == before


@pytest.mark.parametrize("data,filename", [
    (b"a,a\n1,2", "import.csv"), (b"bad", "data.txt"), (b"\xff\xfe", "import.csv"),
    (b"\x00", "import.csv"), (csv_bytes([csv_row(), csv_row()]), "import.csv"),
    (csv_bytes([csv_row(course="UNKNOWN")]), "import.csv"),
])
def test_invalid_csv_rejected_without_changes(app, client, data, filename):
    before = rows(app)
    assert upload(client, data, filename=filename).status_code == 400
    assert rows(app) == before


def test_legacy_csv_is_preserved_without_invented_course_marks(app, client):
    data = b"student_id,name,maths,science,english,attendance\nOLD001,Legacy Student,80,70,60,88.50\n"
    assert upload(client, data).status_code == 302
    student = rows(app, "SELECT * FROM students WHERE student_id='OLD001'")[0]
    page = client.get("/students/" + str(student["id"])).get_data(as_text=True)
    assert "Original three-subject analysis" in page
    records = rows(app, "SELECT * FROM records WHERE student_pk=?", (student["id"],))
    assert all(r["max_hours"] == 0 and r["assignment_1"] is None for r in records)


def test_csv_template_roundtrip_and_filtered_export(app, client):
    template = client.get("/csv/template")
    assert template.mimetype == "text/csv" and b"university_final" in template.data
    assert upload(client, template.data).status_code == 302
    exported = client.get("/csv/export?section=CSE-B&course=21CSC201J")
    result = list(csv.DictReader(io.StringIO(exported.data.decode("utf-8-sig"))))
    assert len(result) == 8
    assert all(r["section"] == "CSE-B" and r["course_code"] == "21CSC201J" for r in result)
    assert upload(client, exported.data, "update").status_code == 302


def test_weighted_summary_and_per_course_risk():
    blank = {key: None for key, _label, _cap in COMPONENTS}
    one = enrich({"student_pk": 1, "max_hours": 100, "attended_hours": 100, **blank})
    two = enrich({"student_pk": 2, "max_hours": 4, "attended_hours": 1, **blank})
    summary = summarize([one, two])
    assert summary["attendance_pct"] > 75 and summary["at_risk"] == 1
    assert summary["max_hours"] == 104 and summary["attended_hours"] == 101 and summary["absent_hours"] == 3
    assert summary["margin"] == one["margin"] + two["margin"]
    assert summary["average"] is None


def test_department_reports_and_exact_student_scope(app, client):
    pdf = client.get("/reports/pdf?section=CSE-A&course=21CSC201J")
    assert pdf.mimetype == "application/pdf"
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.data)).pages)
    normalized = re.sub(r"\s+", " ", text)
    assert "Department Academic Performance Report" in normalized and "Assignment 1" in normalized
    assert "CSE25A001" in text and "CSE25B001" not in text
    workbook = client.get("/reports/xlsx?section=CSE-B&course=21CSC202J")
    wb = load_workbook(io.BytesIO(workbook.data))
    assert wb.sheetnames == ["Report", "Attendance", "Marks"]
    assert wb["Attendance"].max_row == 10  # header + 8 students + summary
    headers = [c.value for c in wb["Marks"][1]]
    assert headers[4:11] == [label + " /" + str(cap) for _key, label, cap in COMPONENTS]
    csv_report = client.get("/reports/csv")
    assert b"Class summary" in csv_report.data
    exact = client.get("/reports/xlsx?student=1")
    wb = load_workbook(io.BytesIO(exact.data))
    assert wb["Marks"].max_row == 8
    assert {wb["Marks"].cell(i, 1).value for i in range(2, 9)} == {"CSE25A001"}
    assert client.get("/reports/not-a-format").status_code == 404


def test_warning_download_snapshot_archive_ownership_and_deleted_student(app, client):
    risk = next(r for r in rows(app) if r["max_hours"] and 4 * r["attended_hours"] < 3 * r["max_hours"])
    response = post(client, "/warnings/issue/" + str(risk["id"]))
    assert response.status_code == 302
    file = client.get(response.location)
    original_text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(file.data)).pages)
    assert "Attendance Shortage Notice" in original_text
    assert risk["course_code"] in original_text
    letter = rows(app, "SELECT * FROM warning_letters")[0]
    snapshot = json.loads(letter["snapshot_json"])
    assert snapshot["record"]["attended_hours"] == risk["attended_hours"]
    assert post(client, "/records/" + str(risk["id"]) + "/attendance", {
        "version": risk["version"], "max_hours": "100", "attended_hours": "100"}).status_code == 302
    assert post(client, "/warnings/issue/" + str(risk["id"])).status_code == 400
    assert post(client, "/students/" + str(risk["student_pk"]) + "/delete").status_code == 302
    later_text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(client.get(response.location).data)).pages)
    assert later_text == original_text
    assert client.get("/warnings").status_code == 200
    second = app.test_client()
    register(second, "other@example.edu")
    second.get("/teacher/dashboard")
    assert second.get(response.location).status_code == 404


def test_settings_profile_and_password(app, client):
    settings = {"action": "profile", "name": "Updated Faculty", "faculty_no": "FAC101",
                "institution": "SRM Institute of Science and Technology", "campus": "Delhi NCR Campus",
                "department": "Computer Science and Engineering", "academic_year": "2026-2027", "semester": "3"}
    assert post(client, "/settings", settings).status_code == 302
    assert b"Updated Faculty" in client.get("/teacher/dashboard").data
    assert post(client, "/settings", {**settings, "academic_year": "2026-2030"}).status_code == 400
    assert post(client, "/settings", {"action": "password", "current_password": "FacultyPass123",
                                    "new_password": "NewFacultyPass456", "confirm_password": "NewFacultyPass456"}).status_code == 302
    client.get("/settings")
    assert check_password_hash(rows(app, "SELECT * FROM teachers")[0]["password_hash"], "NewFacultyPass456")
    post(client, "/logout")
    client.get("/login")
    assert post(client, "/login", {"email": "faculty@example.edu", "password": "NewFacultyPass456"}).status_code == 302


def test_all_pages_links_assets_forms_and_modal_fallbacks(client, app):
    pages = ["/teacher/dashboard", "/students", "/students/add", "/students/1", "/students/1/edit",
             "/students/1/remove", "/attendance", "/attendance/edit", "/marks", "/courses", "/import",
             "/reports", "/warnings", "/settings", "/records/1/attendance", "/records/1/marks",
             "/attendance/session?section=CSE-A&batch=2025-2029"]
    get_rules = {rule.rule for rule in app.url_map.iter_rules() if "GET" in rule.methods}
    post_rules = [rule for rule in app.url_map.iter_rules() if "POST" in rule.methods]
    checked = set()
    for page in pages:
        response = client.get(page)
        assert response.status_code == 200, page
        markup = response.get_data(as_text=True)
        parsed = Links()
        parsed.feed(markup)
        assert parsed.triggers <= parsed.dialogs
        for link in parsed.links + parsed.assets:
            if not link or link.startswith("#"):
                continue
            assert link.startswith("/"), (page, link)
            if link not in checked:
                assert client.get(link).status_code == 200, (page, link)
                checked.add(link)
        for form in parsed.forms:
            assert form.get("method", "get").lower() in ("get", "post")
            action = urlsplit(form.get("action") or page).path
            matched = app.url_map.bind("localhost").match(action, method=form.get("method", "GET").upper())
            assert matched
        assert 'href="#"' not in markup
    assert len(checked) >= 65


def test_upgrade_original_database_backup_and_archive(tmp_path):
    database = tmp_path / "old.sqlite3"
    db = sqlite3.connect(database)
    db.executescript("""
    CREATE TABLE teachers(id INTEGER PRIMARY KEY, name TEXT, email TEXT UNIQUE COLLATE NOCASE,
                          password_hash TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE students(id INTEGER PRIMARY KEY, teacher_id INTEGER REFERENCES teachers(id),
                          student_id TEXT, name TEXT, maths INTEGER, science INTEGER, english INTEGER,
                          attendance REAL, total INTEGER, percentage REAL, grade TEXT, status TEXT);
    """)
    db.execute("INSERT INTO teachers(id,name,email,password_hash) VALUES (1,?,?,?)",
               ("Existing Faculty", "existing@example.edu", generate_password_hash("ExistingPass123")))
    db.execute("INSERT INTO students VALUES (99,1,'OLD099','Existing Student',80,70,60,88,210,70,'B','On Track')")
    db.commit()
    db.close()
    app = create_app({"TESTING": True, "DATABASE": str(database), "SECRET_KEY": "migration-key" * 5})
    students = rows(app, "SELECT * FROM students")
    assert len(students) == 1 and students[0]["id"] == 99
    assert len(rows(app)) == 7
    assert all(r["assignment_1"] is None and r["max_hours"] == 0 for r in rows(app))
    archived = json.loads(rows(app, "SELECT * FROM legacy_results")[0]["source_json"])
    assert archived["maths"] == 80 and archived["percentage"] == 70
    assert len(list(tmp_path.glob("old.sqlite3.v1-backup-*"))) == 1
    client = app.test_client()
    client.get("/login")
    assert post(client, "/login", {"email": "existing@example.edu", "password": "ExistingPass123"}).status_code == 302
    assert client.get("/students/99").status_code == 200
    create_app(dict(app.config))
    assert len(rows(app, "SELECT * FROM students")) == 1


def test_empty_class_all_views_and_exports(tmp_path):
    app = create_app({"TESTING": True, "DATABASE": str(tmp_path / "empty.sqlite3"),
                      "SECRET_KEY": "empty-key" * 8, "SEED_DEMO": False})
    client = app.test_client()
    register(client)
    client.get("/teacher/dashboard")
    for path in ("/teacher/dashboard", "/students", "/attendance", "/attendance/edit", "/marks",
                 "/courses", "/reports", "/reports/pdf", "/reports/xlsx", "/reports/csv", "/warnings"):
        assert client.get(path).status_code == 200, path


def test_exports_neutralize_spreadsheet_formulas(app, client):
    response = post(client, "/students/add", {"student_id": "FORMULA001", "name": "=1+1",
                                            "section": "CSE-A", "batch": "2025-2029"})
    assert response.status_code == 302
    export = client.get("/csv/export?q=FORMULA001")
    assert b"'=1+1" in export.data
    workbook = client.get("/reports/xlsx?q=FORMULA001")
    wb = load_workbook(io.BytesIO(workbook.data))
    assert wb["Attendance"]["B2"].value == "'=1+1"
    assert wb["Attendance"]["B2"].data_type != "f"
